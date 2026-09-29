"""Render the reconciliation results as the review workbook.

This module takes already-prepared results (validation, matching and metrics
output) and turns them into ``outputs/reconciliation_report.xlsx``. It performs
no reconciliation logic of its own: it selects sheets, lays out the Summary and
applies presentation.

Layout (``docs/PLAN.md`` §6)::

    Summary · Matched Records · Missing Payments · Missing Orders ·
    Amount Mismatches · Duplicates · Currency Mismatches ·
    Status Inconsistencies · QA Issues · All Exceptions

Presentation: slate header row with bold white text, frozen header pane,
autofilter on the table tabs, column widths from content, ``#,##0.00`` for money,
``yyyy-mm-dd`` for dates, a percentage for the reconciliation rate, conditional
formatting on tolerance breaches and duplicate/exception flags, and tab colours
grouped by category.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Mapping, Sequence

import pandas as pd
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from recon import __version__
from recon.config import Config
from recon.excel_io import SOURCE_ROW_COLUMN, ReportWorkbookWriter
from recon.matching import (
    DUPLICATE_GROUP_SIZE_COLUMN,
    EXCEPTION_TYPE_COLUMN,
    IS_PRIMARY_COLUMN,
    ReconciliationResult,
)
from recon.models import CHECKS, CHECK_BY_TYPE, ExceptionType, SummaryMetrics
from recon.validation import SOURCE_FILE_COLUMN, RowValidationResult

DEFAULT_CONFIG = Config()

SUMMARY_SHEET = "Summary"
MATCHED_SHEET = "Matched Records"
ALL_EXCEPTIONS_SHEET = "All Exceptions"

#: Exception tabs, in the order declared by :data:`recon.models.CHECKS`.
EXCEPTION_SHEET_NAMES: tuple[str, ...] = tuple(
    dict.fromkeys(check.sheet_name for check in CHECKS)
)
SHEET_ORDER: tuple[str, ...] = (
    SUMMARY_SHEET,
    MATCHED_SHEET,
    *EXCEPTION_SHEET_NAMES,
    ALL_EXCEPTIONS_SHEET,
)

METRIC_COLUMN = "Metric"
VALUE_COLUMN = "Value"

MONEY_NUMBER_FORMAT = "#,##0.00"
DATE_NUMBER_FORMAT = "yyyy-mm-dd"
TIMESTAMP_NUMBER_FORMAT = "yyyy-mm-dd hh:mm:ss"
COUNT_NUMBER_FORMAT = "#,##0"
PERCENT_NUMBER_FORMAT = "0.00%"

MONEY_COLUMNS: tuple[str, ...] = (
    "order_amount",
    "payment_amount",
    "total_paid",
    "difference",
)
DATE_COLUMNS: tuple[str, ...] = ("order_date", "payment_date")
COUNT_COLUMNS: tuple[str, ...] = (
    "source_row",
    "payment_count",
    DUPLICATE_GROUP_SIZE_COLUMN,
)

HEADER_FILL = PatternFill("solid", start_color="FF334155", end_color="FF334155")
HEADER_FONT = Font(bold=True, color="FFFFFFFF")
SECTION_FILL = PatternFill("solid", start_color="FFF1F5F9", end_color="FFF1F5F9")
SECTION_FONT = Font(bold=True, color="FF1E293B")
AMBER_FILL = PatternFill("solid", start_color="FFFEF3C7", end_color="FFFEF3C7")
ALERT_FILL = PatternFill("solid", start_color="FFFEE2E2", end_color="FFFEE2E2")
ALERT_FONT = Font(bold=True, color="FFB91C1C")

#: Tab colours, grouped by what the tab is for.
TAB_COLORS: Mapping[str, str] = {
    SUMMARY_SHEET: "334155",
    MATCHED_SHEET: "15803D",
    "Missing Payments": "B91C1C",
    "Missing Orders": "B91C1C",
    "Amount Mismatches": "D97706",
    "Duplicates": "EA580C",
    "Currency Mismatches": "D97706",
    "Status Inconsistencies": "D97706",
    "QA Issues": "7C3AED",
    ALL_EXCEPTIONS_SHEET: "7F1D1D",
}

MAX_COLUMN_WIDTH = 60
MIN_COLUMN_WIDTH = 10


@dataclass(frozen=True, slots=True)
class _SummaryEntry:
    """One line of the Summary tab."""

    label: str
    value: object
    number_format: str | None = None
    section: bool = False


def write_report(
    path: Path | str,
    *,
    orders: RowValidationResult,
    payments: RowValidationResult,
    reconciliation: ReconciliationResult,
    metrics: SummaryMetrics,
    config: Config = DEFAULT_CONFIG,
    seed: int | None = None,
    generated_at: datetime | None = None,
    overwrite: bool = True,
) -> Path:
    """Write the review workbook and return its path."""

    entries = _summary_entries(
        orders=orders,
        payments=payments,
        metrics=metrics,
        config=config,
        seed=seed,
        generated_at=generated_at or _utc_now(),
    )

    with ReportWorkbookWriter(path, overwrite=overwrite) as writer:
        summary_frame = pd.DataFrame(
            {
                METRIC_COLUMN: [entry.label for entry in entries],
                VALUE_COLUMN: [entry.value for entry in entries],
            }
        )
        summary_sheet = writer.write_sheet(SUMMARY_SHEET, summary_frame)
        _style_summary(summary_sheet, entries)

        tables: Sequence[tuple[str, pd.DataFrame]] = (
            (MATCHED_SHEET, reconciliation.matched),
            (_sheet_name(ExceptionType.MISSING_PAYMENT), reconciliation.missing_payments),
            (_sheet_name(ExceptionType.MISSING_ORDER), reconciliation.missing_orders),
            (_sheet_name(ExceptionType.AMOUNT_MISMATCH), reconciliation.amount_mismatches),
            (_sheet_name(ExceptionType.DUPLICATE_ORDER), reconciliation.duplicates),
            (
                _sheet_name(ExceptionType.CURRENCY_MISMATCH),
                reconciliation.currency_mismatches,
            ),
            (
                _sheet_name(ExceptionType.STATUS_INCONSISTENCY),
                reconciliation.status_inconsistencies,
            ),
            (_sheet_name(ExceptionType.QA_INPUT_ERROR), _qa_issues(orders, payments)),
            (ALL_EXCEPTIONS_SHEET, reconciliation.exceptions),
        )
        for sheet_name, frame in tables:
            worksheet = writer.write_sheet(sheet_name, frame)
            _style_table(worksheet, frame, sheet_name)

    return Path(path)


# --------------------------------------------------------------------------
# Summary
# --------------------------------------------------------------------------


def _utc_now() -> datetime:
    """Naive UTC timestamp; Excel does not accept timezone-aware datetimes."""

    return datetime.now(timezone.utc).replace(tzinfo=None)


def _sheet_name(exception_type: ExceptionType) -> str:
    """Tab name for a check, taken from the model so it cannot drift."""

    return CHECK_BY_TYPE[exception_type].sheet_name


def _summary_entries(
    *,
    orders: RowValidationResult,
    payments: RowValidationResult,
    metrics: SummaryMetrics,
    config: Config,
    seed: int | None,
    generated_at: datetime,
) -> list[_SummaryEntry]:
    entries = [
        _SummaryEntry("Run", None, section=True),
        _SummaryEntry(
            "Report generated (UTC)", generated_at, TIMESTAMP_NUMBER_FORMAT
        ),
        _SummaryEntry("Application version", __version__),
        _SummaryEntry("Orders file", orders.source_file),
        _SummaryEntry("Payments file", payments.source_file),
        _SummaryEntry(
            "Amount tolerance",
            float(config.amount_tolerance),
            MONEY_NUMBER_FORMAT,
        ),
    ]
    if seed is not None:
        entries.append(
            _SummaryEntry("Sample data seed", int(seed), COUNT_NUMBER_FORMAT)
        )

    entries.extend(
        [
            _SummaryEntry("Reconciliation", None, section=True),
            _SummaryEntry("Total orders", metrics.total_orders, COUNT_NUMBER_FORMAT),
            _SummaryEntry(
                "Total payments", metrics.total_payments, COUNT_NUMBER_FORMAT
            ),
            _SummaryEntry(
                "Matched records", metrics.matched_count, COUNT_NUMBER_FORMAT
            ),
            _SummaryEntry(
                "Exception count (affected records)",
                metrics.exception_count,
                COUNT_NUMBER_FORMAT,
            ),
            _SummaryEntry(
                "Quarantined rows (QA issues)",
                metrics.quarantined_rows,
                COUNT_NUMBER_FORMAT,
            ),
            _SummaryEntry(
                "Reconciliation rate",
                float(metrics.reconciliation_rate),
                PERCENT_NUMBER_FORMAT,
            ),
            _SummaryEntry("Exception findings by type", None, section=True),
        ]
    )
    for check in CHECKS:
        label = str(check.exception_type)
        entries.append(
            _SummaryEntry(
                label, int(metrics.exception_counts.get(label, 0)), COUNT_NUMBER_FORMAT
            )
        )
    return entries


def _style_summary(worksheet: Worksheet, entries: Sequence[_SummaryEntry]) -> None:
    _style_header(worksheet)
    worksheet.freeze_panes = "A2"
    worksheet.sheet_properties.tabColor = f"FF{TAB_COLORS[SUMMARY_SHEET]}"
    worksheet.column_dimensions["A"].width = 38
    worksheet.column_dimensions["B"].width = 34

    for offset, entry in enumerate(entries, start=2):
        label_cell = worksheet.cell(row=offset, column=1)
        value_cell = worksheet.cell(row=offset, column=2)
        if entry.section:
            for cell in (label_cell, value_cell):
                cell.fill = SECTION_FILL
                cell.font = SECTION_FONT
            continue
        label_cell.alignment = Alignment(indent=1)
        if entry.number_format:
            value_cell.number_format = entry.number_format


# --------------------------------------------------------------------------
# Tables
# --------------------------------------------------------------------------


def _qa_issues(
    orders: RowValidationResult, payments: RowValidationResult
) -> pd.DataFrame:
    """Both quarantine frames in one sheet, ordered by file and row."""

    frames = [
        frame
        for frame in (orders.quarantined, payments.quarantined)
        if not frame.empty
    ]
    if not frames:
        columns = list(
            dict.fromkeys(
                [*orders.quarantined.columns, *payments.quarantined.columns]
            )
        )
        return pd.DataFrame(columns=columns)
    combined = pd.concat(frames, ignore_index=True)
    return combined.sort_values(
        [SOURCE_FILE_COLUMN, SOURCE_ROW_COLUMN], kind="stable", ignore_index=True
    )


def _style_table(worksheet: Worksheet, frame: pd.DataFrame, sheet_name: str) -> None:
    _style_header(worksheet)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_properties.tabColor = f"FF{TAB_COLORS.get(sheet_name, '334155')}"

    for column, width in _column_widths(frame).items():
        worksheet.column_dimensions[column].width = width
    for column, number_format in _number_formats(frame).items():
        letter = _column_letter(frame, column)
        for row in range(2, len(frame) + 2):
            worksheet[f"{letter}{row}"].number_format = number_format

    _add_conditional_formats(worksheet, frame, sheet_name)


def _style_header(worksheet: Worksheet) -> None:
    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center")
    worksheet.row_dimensions[1].height = 18


def _column_widths(frame: pd.DataFrame) -> dict[str, int]:
    """Column letter -> width, sized from the header and the longest value."""

    widths: dict[str, int] = {}
    for position, column in enumerate(frame.columns, start=1):
        longest = len(str(column))
        if not frame.empty:
            lengths = frame[column].astype("string").str.len()
            if lengths.notna().any():
                longest = max(longest, int(lengths.max()))
        widths[get_column_letter(position)] = max(
            MIN_COLUMN_WIDTH, min(MAX_COLUMN_WIDTH, longest + 2)
        )
    return widths


def _number_formats(frame: pd.DataFrame) -> dict[str, str]:
    formats: dict[str, str] = {}
    for column in frame.columns:
        name = str(column)
        if name in MONEY_COLUMNS:
            formats[name] = MONEY_NUMBER_FORMAT
        elif name in DATE_COLUMNS:
            formats[name] = DATE_NUMBER_FORMAT
        elif name in COUNT_COLUMNS:
            formats[name] = COUNT_NUMBER_FORMAT
    return formats


def _column_letter(frame: pd.DataFrame, column: str) -> str:
    return get_column_letter(list(frame.columns).index(column) + 1)


def _add_conditional_formats(
    worksheet: Worksheet, frame: pd.DataFrame, sheet_name: str
) -> None:
    if frame.empty:
        return

    last_row = len(frame) + 1
    if sheet_name == "Amount Mismatches" and "difference" in frame.columns:
        # bigger breaches stand out more
        letter = _column_letter(frame, "difference")
        worksheet.conditional_formatting.add(
            f"{letter}2:{letter}{last_row}",
            ColorScaleRule(
                start_type="min",
                start_color="FFFFFFFF",
                mid_type="percentile",
                mid_value=50,
                mid_color="FFFDE68A",
                end_type="max",
                end_color="FFF87171",
            ),
        )

    if sheet_name == "Duplicates" and IS_PRIMARY_COLUMN in frame.columns:
        letter = _column_letter(frame, IS_PRIMARY_COLUMN)
        worksheet.conditional_formatting.add(
            f"A2:{get_column_letter(len(frame.columns))}{last_row}",
            FormulaRule(formula=[f'${letter}2=FALSE'], fill=AMBER_FILL),
        )

    if sheet_name == ALL_EXCEPTIONS_SHEET and EXCEPTION_TYPE_COLUMN in frame.columns:
        letter = _column_letter(frame, EXCEPTION_TYPE_COLUMN)
        worksheet.conditional_formatting.add(
            f"{letter}2:{letter}{last_row}",
            FormulaRule(
                formula=[
                    f'OR(${letter}2="{ExceptionType.MISSING_PAYMENT}",'
                    f'${letter}2="{ExceptionType.MISSING_ORDER}")'
                ],
                font=ALERT_FONT,
                fill=ALERT_FILL,
            ),
        )
