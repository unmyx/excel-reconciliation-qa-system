"""End-to-end tests: small exports -> pipeline -> workbook, plus CLI behaviour.

Everything runs in ``tmp_path``; no test depends on the repository's generated
sample data.
"""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

import run as cli
from recon.config import Config
from recon.excel_io import ExcelInputError
from recon.matching import IS_PRIMARY_COLUMN
from recon.pipeline import run_pipeline
from recon.report import SHEET_ORDER, SUMMARY_SHEET
from recon.validation import InputValidationError

ORDER_COLUMNS = [
    "order_id",
    "customer_id",
    "order_date",
    "currency",
    "order_amount",
    "order_status",
    "sales_channel",
]
PAYMENT_COLUMNS = [
    "payment_id",
    "order_id",
    "payment_date",
    "currency",
    "payment_amount",
    "payment_status",
    "payment_method",
]


def _orders_frame() -> pd.DataFrame:
    return pd.DataFrame(
        [
            # matched 1:1
            ["ORD-000001", "CUST-00001", "2025-01-02", "EUR", 100.0, "completed", "web"],
            # no payment at all
            ["ORD-000002", "CUST-00002", "2025-01-03", "EUR", 50.0, "completed", "web"],
            # paid 70 of 80
            ["ORD-000003", "CUST-00003", "2025-01-04", "EUR", 80.0, "completed", "web"],
            # duplicated order row (both rows identical)
            ["ORD-000004", "CUST-00004", "2025-01-05", "EUR", 20.0, "completed", "web"],
            ["ORD-000004", "CUST-00004", "2025-01-05", "EUR", 20.0, "completed", "web"],
            # data-quality row: currency missing
            ["ORD-000006", "CUST-00006", "2025-01-07", None, 10.0, "completed", "web"],
        ],
        columns=ORDER_COLUMNS,
    )


def _payments_frame() -> pd.DataFrame:
    return pd.DataFrame(
        [
            ["PAY-0000001", "ORD-000001", "2025-01-03", "EUR", 100.0, "successful", "card"],
            ["PAY-0000003", "ORD-000003", "2025-01-05", "EUR", 70.0, "successful", "card"],
            ["PAY-0000004", "ORD-000004", "2025-01-06", "EUR", 20.0, "successful", "card"],
            # order id that exists nowhere in the orders export
            ["PAY-0000007", "ORD-000099", "2025-01-08", "EUR", 10.0, "successful", "card"],
        ],
        columns=PAYMENT_COLUMNS,
    )


def _write_export(path: Path, frame: pd.DataFrame, sheet_name: str) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    frame.to_excel(path, sheet_name=sheet_name, index=False)
    return path


def _config(tmp_path: Path, **overrides: object) -> Config:
    settings: dict[str, object] = {
        "orders_path": tmp_path / "data" / "orders_export.xlsx",
        "payments_path": tmp_path / "data" / "payments_export.xlsx",
        "output_path": tmp_path / "outputs" / "reconciliation_report.xlsx",
        "manifest_path": tmp_path / "data" / "generation_manifest.json",
    }
    settings.update(overrides)
    return Config(**settings)


def _prepare(tmp_path: Path, **config_overrides: object) -> Config:
    config = _config(tmp_path, **config_overrides)
    _write_export(Path(config.orders_path), _orders_frame(), "orders")
    _write_export(Path(config.payments_path), _payments_frame(), "payments")
    return config


def _write_manifest(
    config: Config, *, seed: int = 42, orders_file: str = "data/orders_export.xlsx"
) -> None:
    Path(config.manifest_path).write_text(
        json.dumps({"seed": seed, "outputs": {"orders_file": orders_file}}),
        encoding="utf-8",
    )


def _summary_values(path: Path) -> dict[str, object]:
    worksheet = load_workbook(path)[SUMMARY_SHEET]
    return {
        row[0].value: row[1].value
        for row in worksheet.iter_rows(min_row=2)
        if row[0].value
    }


# --------------------------------------------------------------------------
# Pipeline
# --------------------------------------------------------------------------


def test_pipeline_runs_every_stage_and_writes_the_report(tmp_path: Path) -> None:
    config = _prepare(tmp_path)

    result = run_pipeline(config)

    assert result.output_path == Path(config.output_path)
    assert result.output_path.is_file()
    workbook = load_workbook(result.output_path)
    assert workbook.sheetnames == list(SHEET_ORDER)
    assert result.orders_file == "orders_export.xlsx"
    assert result.payments_file == "payments_export.xlsx"


def test_pipeline_metrics_match_the_exports(tmp_path: Path) -> None:
    config = _prepare(tmp_path)

    metrics = run_pipeline(config).metrics

    # rows read before quarantine
    assert metrics.total_orders == len(_orders_frame()) == 6
    assert metrics.total_payments == len(_payments_frame()) == 4
    # ORD-000001 and the primary row of ORD-000004 match
    assert metrics.matched_count == 2
    # missing payment, amount mismatch, missing order, duplicate order row
    assert metrics.exception_count == 4
    assert metrics.quarantined_rows == 1
    assert metrics.reconciliation_rate == pytest.approx(2 / 6)
    assert metrics.exception_counts == {
        "Missing Payment": 1,
        "Missing Order": 1,
        "Amount Mismatch": 1,
        "Duplicate Order": 1,
        "Input Data Quality Error": 1,
    }


def test_summary_sheet_matches_the_returned_metrics(tmp_path: Path) -> None:
    config = _prepare(tmp_path)

    result = run_pipeline(config)
    values = _summary_values(result.output_path)

    assert values["Total orders"] == result.metrics.total_orders
    assert values["Total payments"] == result.metrics.total_payments
    assert values["Matched records"] == result.metrics.matched_count
    assert values["Exception count (affected records)"] == result.metrics.exception_count
    assert values["Quarantined rows (QA issues)"] == result.metrics.quarantined_rows
    assert values["Reconciliation rate"] == pytest.approx(
        result.metrics.reconciliation_rate
    )
    assert values["Orders file"] == "orders_export.xlsx"


def test_pipeline_preserves_qa_rows_and_their_source_rows(tmp_path: Path) -> None:
    config = _prepare(tmp_path)

    result = run_pipeline(config)
    quality = pd.read_excel(result.output_path, sheet_name="QA Issues")

    assert quality["source_file"].tolist() == ["orders_export.xlsx"]
    assert quality["source_row"].tolist() == [7]  # 6th data row of the orders sheet
    assert quality["qa_reason"].tolist() == ["missing_required_value"]
    assert "currency" in quality["qa_details"].tolist()[0]


def test_pipeline_creates_the_output_directory(tmp_path: Path) -> None:
    config = _prepare(tmp_path)
    assert not Path(config.output_path).parent.exists()

    result = run_pipeline(config)

    assert Path(config.output_path).parent.is_dir()
    assert result.output_path.is_file()


# --------------------------------------------------------------------------
# Seed recorded from the sample-data manifest
# --------------------------------------------------------------------------


def test_seed_is_recorded_when_the_manifest_matches(tmp_path: Path) -> None:
    config = _prepare(tmp_path)
    _write_manifest(config, seed=42)

    result = run_pipeline(config)

    assert _summary_values(result.output_path)["Sample data seed"] == 42


def test_explicit_seed_wins_over_the_manifest(tmp_path: Path) -> None:
    config = _prepare(tmp_path)
    _write_manifest(config, seed=42)

    result = run_pipeline(config, seed=7)

    assert _summary_values(result.output_path)["Sample data seed"] == 7


def test_seed_is_skipped_when_the_manifest_describes_other_files(
    tmp_path: Path,
) -> None:
    config = _prepare(tmp_path)
    _write_manifest(config, seed=42, orders_file="data/somewhere_else.xlsx")

    result = run_pipeline(config)

    assert "Sample data seed" not in _summary_values(result.output_path)


# --------------------------------------------------------------------------
# Fatal input errors write nothing
# --------------------------------------------------------------------------


def test_missing_required_column_stops_before_writing_anything(tmp_path: Path) -> None:
    config = _config(tmp_path)
    _write_export(
        Path(config.orders_path), _orders_frame().drop(columns=["currency"]), "orders"
    )
    _write_export(Path(config.payments_path), _payments_frame(), "payments")

    with pytest.raises(InputValidationError, match="currency"):
        run_pipeline(config)

    assert not Path(config.output_path).exists()


def test_missing_input_file_stops_before_writing_anything(tmp_path: Path) -> None:
    config = _prepare(tmp_path)
    Path(config.payments_path).unlink()

    with pytest.raises(ExcelInputError, match="not found"):
        run_pipeline(config)

    assert not Path(config.output_path).exists()


# --------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------


def test_cli_runs_with_explicit_paths(
    tmp_path: Path, capsys: pytest.CaptureFixture
) -> None:
    config = _prepare(tmp_path)

    exit_code = cli.main(
        [
            "--orders",
            str(config.orders_path),
            "--payments",
            str(config.payments_path),
            "--output",
            str(config.output_path),
        ]
    )

    captured = capsys.readouterr()
    assert exit_code == cli.EXIT_SUCCESS == 0
    assert Path(config.output_path).is_file()
    # logs go to stderr, but nothing should look like a failure
    assert "error:" not in captured.err


def test_cli_defaults_use_the_configured_paths(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    config = _prepare(tmp_path)
    monkeypatch.setattr(cli, "DEFAULT_CONFIG", config)

    exit_code = cli.main([])

    assert exit_code == 0
    assert Path(config.output_path).is_file()


def test_cli_returns_two_for_a_missing_input_file(
    tmp_path: Path, capsys: pytest.CaptureFixture
) -> None:
    config = _config(tmp_path)
    _write_export(Path(config.orders_path), _orders_frame(), "orders")

    exit_code = cli.main(
        [
            "--orders",
            str(config.orders_path),
            "--payments",
            str(config.payments_path),
            "--output",
            str(config.output_path),
        ]
    )

    captured = capsys.readouterr()
    assert exit_code == cli.EXIT_FATAL_INPUT_ERROR == 2
    assert "error:" in captured.err and "not found" in captured.err
    assert not Path(config.output_path).exists()


def test_cli_returns_two_for_a_missing_required_column(
    tmp_path: Path, capsys: pytest.CaptureFixture
) -> None:
    config = _config(tmp_path)
    _write_export(
        Path(config.orders_path), _orders_frame().drop(columns=["order_id"]), "orders"
    )
    _write_export(Path(config.payments_path), _payments_frame(), "payments")

    exit_code = cli.main(
        [
            "--orders",
            str(config.orders_path),
            "--payments",
            str(config.payments_path),
            "--output",
            str(config.output_path),
        ]
    )

    captured = capsys.readouterr()
    assert exit_code == 2
    assert "order_id" in captured.err


def test_cli_returns_two_for_a_negative_tolerance(
    tmp_path: Path, capsys: pytest.CaptureFixture
) -> None:
    config = _prepare(tmp_path)

    exit_code = cli.main(["--tolerance", "-1", "--output", str(config.output_path)])

    assert exit_code == 2
    assert "tolerance" in capsys.readouterr().err


def test_cli_returns_one_for_an_unexpected_error(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture
) -> None:
    def explode(*args: object, **kwargs: object) -> None:
        raise RuntimeError("boom")

    monkeypatch.setattr(cli, "run_pipeline", explode)

    exit_code = cli.main([])

    captured = capsys.readouterr()
    assert exit_code == cli.EXIT_UNEXPECTED_ERROR == 1
    assert "unexpected failure" in captured.err and "RuntimeError" in captured.err


def test_cli_tolerance_override_reaches_the_config(tmp_path: Path) -> None:
    config = _prepare(tmp_path)

    exit_code = cli.main(
        [
            "--orders",
            str(config.orders_path),
            "--payments",
            str(config.payments_path),
            "--output",
            str(config.output_path),
            "--tolerance",
            "0.5",
        ]
    )

    assert exit_code == 0
    assert _summary_values(Path(config.output_path))["Amount tolerance"] == (
        pytest.approx(0.5)
    )


def test_cli_generate_sample_builds_exports_then_reconciles(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    config = _config(tmp_path, order_count=800)
    monkeypatch.setattr(cli, "DEFAULT_CONFIG", config)

    exit_code = cli.main(["--generate-sample", "--seed", "7"])

    assert exit_code == 0
    assert Path(config.orders_path).is_file()
    assert Path(config.payments_path).is_file()
    assert Path(config.output_path).is_file()

    values = _summary_values(Path(config.output_path))
    assert values["Sample data seed"] == 7
    assert values["Total orders"] > 800  # 800 orders plus the injected duplicates
    assert load_workbook(Path(config.output_path)).sheetnames == list(SHEET_ORDER)


def test_build_config_applies_overrides(tmp_path: Path) -> None:
    args = cli.parse_args(
        [
            "--orders",
            str(tmp_path / "o.xlsx"),
            "--payments",
            str(tmp_path / "p.xlsx"),
            "--output",
            str(tmp_path / "r.xlsx"),
            "--tolerance",
            "0.25",
        ]
    )

    config = cli.build_config(args)

    assert config.orders_path == tmp_path / "o.xlsx"
    assert config.payments_path == tmp_path / "p.xlsx"
    assert config.output_path == tmp_path / "r.xlsx"
    assert config.amount_tolerance == pytest.approx(0.25)
    assert config.expected_currencies == cli.DEFAULT_CONFIG.expected_currencies


# --------------------------------------------------------------------------
# The manifest is the ground truth for a seeded sample run
# --------------------------------------------------------------------------


def test_pipeline_finds_exactly_what_the_generator_injected(tmp_path: Path) -> None:
    """Every injected anomaly must be detected, cascades included."""

    from dataclasses import replace

    from scripts.generate_sample_data import InjectionCounts, generate_dataset

    counts = InjectionCounts(
        missing_payment=12,
        missing_order=6,
        amount_mismatch=9,
        currency_mismatch=5,
        status_cancelled_with_successful_payment=3,
        status_completed_with_failed_payment=2,
        duplicate_order=3,
        duplicate_payment_same_id=3,
        duplicate_payment_new_id=2,
        qa_order_rows=10,
        qa_payment_rows=10,
    )
    config = _config(tmp_path)
    generated = generate_dataset(
        order_count=800,
        seed=42,
        output_dir=Path(config.orders_path).parent,
        counts=counts,
    )
    expected = generated.manifest["expected_detections"]

    result = run_pipeline(
        replace(
            config,
            orders_path=generated.orders_path,
            payments_path=generated.payments_path,
            manifest_path=generated.manifest_path,
        )
    )

    reconciliation = result.reconciliation
    assert result.metrics.matched_count == expected["expected_matched_orders"]
    assert len(reconciliation.missing_payments) == expected["missing_payments"]
    assert len(reconciliation.missing_orders) == expected["missing_orders"]
    assert len(reconciliation.amount_mismatches) == expected["amount_mismatches"]
    assert len(reconciliation.currency_mismatches) == expected["currency_mismatches"]
    assert (
        len(reconciliation.status_inconsistencies)
        == expected["status_inconsistencies"]
    )
    assert (
        int((~reconciliation.duplicate_orders[IS_PRIMARY_COLUMN]).sum())
        == expected["duplicate_order_findings"]
    )
    assert (
        int((~reconciliation.duplicate_payments[IS_PRIMARY_COLUMN]).sum())
        == expected["duplicate_payment_findings"]
    )
    assert (
        len(reconciliation.blocked_payments)
        == expected["payments_blocked_by_quarantined_orders"]
    )

    # quarantine accounting, by source and by reason
    assert result.metrics.quarantined_rows == expected["quarantined_rows"]
    assert result.orders.quarantined_count == (
        expected["quarantined_rows_by_source"]["orders"]
    )
    assert result.payments.quarantined_count == (
        expected["quarantined_rows_by_source"]["payments"]
    )
    reasons: dict[str, int] = {}
    for validation in (result.orders, result.payments):
        for reason, count in validation.reason_counts.items():
            reasons[reason] = reasons.get(reason, 0) + count
    assert reasons == expected["quarantined_rows_by_reason"]

    # the union holds one finding per (record, exception type)
    assert reconciliation.finding_count == sum(
        result.metrics.exception_counts.values()
    )
    assert result.metrics.exception_counts["Input Data Quality Error"] == (
        expected["input_data_quality_findings"]
    )
