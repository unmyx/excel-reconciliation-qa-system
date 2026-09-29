"""Tests for the review workbook: sheet set, values and presentation."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from conftest import (
    ORDERS_FILE,
    PAYMENTS_FILE,
    order_row,
    orders_result,
    payment_row,
    payments_result,
    quarantined_row,
)
from recon.excel_io import SOURCE_ROW_COLUMN
from recon.matching import (
    DUPLICATE_GROUP_SIZE_COLUMN,
    EXCEPTION_TYPE_COLUMN,
    IS_PRIMARY_COLUMN,
    reconcile,
)
from recon.metrics import summarize
from recon.models import CHECKS, ExceptionType
from recon.report import (
    ALL_EXCEPTIONS_SHEET,
    DATE_NUMBER_FORMAT,
    MONEY_NUMBER_FORMAT,
    PERCENT_NUMBER_FORMAT,
    SHEET_ORDER,
    SUMMARY_SHEET,
    TAB_COLORS,
    write_report,
)
from recon.validation import QA_REASON_COLUMN

GENERATED_AT = datetime(2025, 6, 1, 12, 30, 0)


def _run(
    orders: list[dict[str, object]],
    payments: list[dict[str, object]],
    *,
    quarantined_orders: list[dict[str, object]] | None = None,
    quarantined_payments: list[dict[str, object]] | None = None,
) -> tuple:
    checked_orders = orders_result(orders, quarantined=quarantined_orders or [])
    checked_payments = payments_result(payments, quarantined=quarantined_payments or [])
    reconciliation = reconcile(checked_orders, checked_payments)
    metrics = summarize(checked_orders, checked_payments, reconciliation)
    return checked_orders, checked_payments, reconciliation, metrics


def _write(path: Path, run: tuple, **kwargs: object) -> Path:
    checked_orders, checked_payments, reconciliation, metrics = run
    return write_report(
        path,
        orders=checked_orders,
        payments=checked_payments,
        reconciliation=reconciliation,
        metrics=metrics,
        generated_at=GENERATED_AT,
        **kwargs,
    )


@pytest.fixture
def sample_run() -> tuple:
    """One run covering matched rows, exceptions, duplicates and QA rows."""

    orders = [
        order_row("ORD-000001", 100.0, source_row=2),  # matched
        order_row("ORD-000002", 50.0, source_row=3),  # missing payment
        order_row("ORD-000003", 80.0, source_row=4),  # amount mismatch
        order_row("ORD-000004", 20.0, source_row=5),  # duplicate order row
        order_row("ORD-000004", 20.0, source_row=6),
        order_row("ORD-000005", 30.0, source_row=7, currency="EUR"),  # currency mismatch
        order_row("ORD-000006", 40.0, source_row=8, status="CANCELLED"),  # status
    ]
    payments = [
        payment_row("PAY-0000001", "ORD-000001", 100.0, source_row=2),
        payment_row("PAY-0000003", "ORD-000003", 70.0, source_row=3),
        payment_row("PAY-0000004", "ORD-000004", 20.0, source_row=4),
        payment_row("PAY-0000005", "ORD-000005", 30.0, source_row=5, currency="USD"),
        payment_row(
            "PAY-0000006", "ORD-000006", 40.0, source_row=6, status="SUCCESSFUL"
        ),
        payment_row("PAY-0000007", "ORD-000099", 10.0, source_row=7),  # missing order
    ]
    return _run(
        orders,
        payments,
        quarantined_orders=[
            quarantined_row(
                source_row=11,
                source_file=ORDERS_FILE,
                order_id="ORD-000011",
                reason="malformed_id",
                details="order_id: malformed_id",
            )
        ],
        quarantined_payments=[
            quarantined_row(
                source_row=12,
                source_file=PAYMENTS_FILE,
                order_id="ORD-000001",
                payment_id="PAY-0000012",
                reason="invalid_numeric",
                details="payment_amount: invalid_numeric",
            )
        ],
    )


# --------------------------------------------------------------------------
# Workbook shape
# --------------------------------------------------------------------------


def test_sheet_order_is_exactly_as_planned(tmp_path: Path, sample_run: tuple) -> None:
    path = _write(tmp_path / "report.xlsx", sample_run)

    workbook = load_workbook(path)

    assert workbook.sheetnames == list(SHEET_ORDER)
    assert workbook.sheetnames == [
        "Summary",
        "Matched Records",
        "Missing Payments",
        "Missing Orders",
        "Amount Mismatches",
        "Duplicates",
        "Currency Mismatches",
        "Status Inconsistencies",
        "QA Issues",
        "All Exceptions",
    ]


def test_required_sheets_all_exist(tmp_path: Path, sample_run: tuple) -> None:
    path = _write(tmp_path / "report.xlsx", sample_run)

    workbook = load_workbook(path)

    assert SUMMARY_SHEET in workbook.sheetnames
    assert ALL_EXCEPTIONS_SHEET in workbook.sheetnames
    for check in CHECKS:
        assert check.sheet_name in workbook.sheetnames


def test_header_row_is_bold_white_on_slate(
    tmp_path: Path, sample_run: tuple
) -> None:
    path = _write(tmp_path / "report.xlsx", sample_run)
    workbook = load_workbook(path)

    for sheet_name in SHEET_ORDER:
        header = workbook[sheet_name][1]
        assert all(cell.font.bold for cell in header)
        assert all(cell.font.color.rgb == "FFFFFFFF" for cell in header)
        assert all(cell.fill.start_color.rgb == "FF334155" for cell in header)

    assert workbook["Matched Records"]["A1"].value == "source_file"
    assert workbook["Summary"]["A1"].value == "Metric"
    assert workbook["Summary"]["B1"].value == "Value"


def test_every_sheet_freezes_its_header_row(
    tmp_path: Path, sample_run: tuple
) -> None:
    path = _write(tmp_path / "report.xlsx", sample_run)
    workbook = load_workbook(path)

    assert all(workbook[name].freeze_panes == "A2" for name in SHEET_ORDER)


def test_table_sheets_have_autofilters(tmp_path: Path, sample_run: tuple) -> None:
    path = _write(tmp_path / "report.xlsx", sample_run)
    workbook = load_workbook(path)

    assert workbook[SUMMARY_SHEET].auto_filter.ref is None  # key/value sheet, not a table
    for sheet_name in SHEET_ORDER:
        if sheet_name == SUMMARY_SHEET:
            continue
        worksheet = workbook[sheet_name]
        assert worksheet.auto_filter.ref is not None
        assert worksheet.auto_filter.ref.startswith("A1")


def test_tab_colours_are_grouped_by_category(tmp_path: Path, sample_run: tuple) -> None:
    path = _write(tmp_path / "report.xlsx", sample_run)
    workbook = load_workbook(path)

    for sheet_name in SHEET_ORDER:
        assert workbook[sheet_name].sheet_properties.tabColor.rgb == (
            f"FF{TAB_COLORS[sheet_name]}"
        )


# --------------------------------------------------------------------------
# Number and date formats
# --------------------------------------------------------------------------


def test_monetary_columns_use_two_decimal_format(
    tmp_path: Path, sample_run: tuple
) -> None:
    path = _write(tmp_path / "report.xlsx", sample_run)
    worksheet = load_workbook(path)["Amount Mismatches"]

    headers = [cell.value for cell in worksheet[1]]
    for column in ("order_amount", "total_paid", "difference"):
        cell = worksheet.cell(row=2, column=headers.index(column) + 1)
        assert cell.number_format == MONEY_NUMBER_FORMAT


def test_date_columns_use_iso_date_format(tmp_path: Path, sample_run: tuple) -> None:
    path = _write(tmp_path / "report.xlsx", sample_run)
    worksheet = load_workbook(path)["Matched Records"]

    headers = [cell.value for cell in worksheet[1]]
    cell = worksheet.cell(row=2, column=headers.index("order_date") + 1)

    assert cell.number_format == DATE_NUMBER_FORMAT
    assert isinstance(cell.value, datetime)


def test_reconciliation_rate_uses_percent_format(
    tmp_path: Path, sample_run: tuple
) -> None:
    path = _write(tmp_path / "report.xlsx", sample_run)
    worksheet = load_workbook(path)[SUMMARY_SHEET]

    rate_row = next(
        row
        for row in worksheet.iter_rows(min_row=2)
        if row[0].value == "Reconciliation rate"
    )

    assert rate_row[1].number_format == PERCENT_NUMBER_FORMAT
    assert 0 < rate_row[1].value <= 1


def test_conditional_formatting_is_applied_where_useful(
    tmp_path: Path, sample_run: tuple
) -> None:
    path = _write(tmp_path / "report.xlsx", sample_run)
    workbook = load_workbook(path)

    assert list(workbook["Amount Mismatches"].conditional_formatting)
    assert list(workbook["Duplicates"].conditional_formatting)
    assert list(workbook[ALL_EXCEPTIONS_SHEET].conditional_formatting)


# --------------------------------------------------------------------------
# Sheet content
# --------------------------------------------------------------------------


def test_summary_reports_run_metadata_and_metrics(
    tmp_path: Path, sample_run: tuple
) -> None:
    path = _write(tmp_path / "report.xlsx", sample_run, seed=42)
    worksheet = load_workbook(path)[SUMMARY_SHEET]

    values = {
        row[0].value: row[1].value
        for row in worksheet.iter_rows(min_row=2)
        if row[0].value
    }

    assert values["Report generated (UTC)"] == GENERATED_AT
    assert values["Orders file"] == ORDERS_FILE
    assert values["Payments file"] == PAYMENTS_FILE
    assert values["Amount tolerance"] == pytest.approx(0.01)
    assert values["Sample data seed"] == 42

    _orders, _payments, _reconciliation, metrics = sample_run
    assert values["Total orders"] == metrics.total_orders
    assert values["Total payments"] == metrics.total_payments
    assert values["Matched records"] == metrics.matched_count
    assert values["Exception count (affected records)"] == metrics.exception_count
    assert values["Quarantined rows (QA issues)"] == metrics.quarantined_rows
    assert values["Missing Payment"] == metrics.exception_counts["Missing Payment"]
    assert values["Input Data Quality Error"] > 0


def test_summary_omits_the_seed_when_it_is_not_known(
    tmp_path: Path, sample_run: tuple
) -> None:
    path = _write(tmp_path / "report.xlsx", sample_run)
    worksheet = load_workbook(path)[SUMMARY_SHEET]

    labels = {row[0].value for row in worksheet.iter_rows(min_row=2)}

    assert "Sample data seed" not in labels


def test_matched_and_exception_sheets_carry_their_rows(
    tmp_path: Path, sample_run: tuple
) -> None:
    path = _write(tmp_path / "report.xlsx", sample_run)

    matched = pd.read_excel(path, sheet_name="Matched Records")
    missing_payments = pd.read_excel(path, sheet_name="Missing Payments")
    missing_orders = pd.read_excel(path, sheet_name="Missing Orders")
    amount_mismatches = pd.read_excel(path, sheet_name="Amount Mismatches")

    # ORD-000004 matches on its primary row even though the row itself is duplicated
    assert matched["order_id"].tolist() == ["ORD-000001", "ORD-000004"]
    assert missing_payments["order_id"].tolist() == ["ORD-000002"]
    assert missing_orders["order_id"].tolist() == ["ORD-000099"]
    assert amount_mismatches["order_id"].tolist() == ["ORD-000003"]
    assert amount_mismatches["difference"].tolist() == [pytest.approx(10.0)]


def test_duplicates_sheet_keeps_duplicate_metadata(
    tmp_path: Path, sample_run: tuple
) -> None:
    path = _write(tmp_path / "report.xlsx", sample_run)

    duplicates = pd.read_excel(path, sheet_name="Duplicates")

    assert "duplicate_type" in duplicates.columns
    assert DUPLICATE_GROUP_SIZE_COLUMN in duplicates.columns
    assert IS_PRIMARY_COLUMN in duplicates.columns
    assert set(duplicates["duplicate_type"]) == {"duplicate_order_id"}
    assert duplicates[DUPLICATE_GROUP_SIZE_COLUMN].tolist() == [2, 2]
    assert duplicates[IS_PRIMARY_COLUMN].tolist() == [True, False]


def test_qa_issues_sheet_keeps_traceability_columns(
    tmp_path: Path, sample_run: tuple
) -> None:
    path = _write(tmp_path / "report.xlsx", sample_run)

    quality = pd.read_excel(path, sheet_name="QA Issues")

    assert list(quality.columns)[:4] == [
        "source_file",
        SOURCE_ROW_COLUMN,
        QA_REASON_COLUMN,
        "qa_details",
    ]
    assert quality["source_file"].tolist() == [ORDERS_FILE, PAYMENTS_FILE]
    assert quality[SOURCE_ROW_COLUMN].tolist() == [11, 12]
    assert quality[QA_REASON_COLUMN].tolist() == ["malformed_id", "invalid_numeric"]


def test_all_exceptions_is_sorted_by_tab_order(
    tmp_path: Path, sample_run: tuple
) -> None:
    path = _write(tmp_path / "report.xlsx", sample_run)

    exceptions = pd.read_excel(path, sheet_name=ALL_EXCEPTIONS_SHEET)

    assert EXCEPTION_TYPE_COLUMN in exceptions.columns
    types = exceptions[EXCEPTION_TYPE_COLUMN].tolist()
    tab_order = {str(check.exception_type): index for index, check in enumerate(CHECKS)}
    assert types == sorted(types, key=lambda label: tab_order[label])


def test_empty_reconciliation_still_writes_every_sheet(tmp_path: Path) -> None:
    run = _run([], [])
    path = _write(tmp_path / "report.xlsx", run)

    workbook = load_workbook(path)

    assert workbook.sheetnames == list(SHEET_ORDER)
    for sheet_name in SHEET_ORDER:
        if sheet_name == SUMMARY_SHEET:
            assert workbook[sheet_name].max_row > 1  # run metadata still present
        else:
            assert workbook[sheet_name].max_row == 1  # header row only
    assert workbook["Matched Records"]["A1"].value == "source_file"


def test_workbook_reopens_with_expected_content(
    tmp_path: Path, sample_run: tuple
) -> None:
    path = _write(tmp_path / "report.xlsx", sample_run)

    reopened = load_workbook(path)
    exceptions = pd.read_excel(path, sheet_name=ALL_EXCEPTIONS_SHEET)

    assert reopened.sheetnames == list(SHEET_ORDER)
    assert set(exceptions[EXCEPTION_TYPE_COLUMN]) >= {
        str(ExceptionType.MISSING_PAYMENT),
        str(ExceptionType.MISSING_ORDER),
        str(ExceptionType.AMOUNT_MISMATCH),
        str(ExceptionType.DUPLICATE_ORDER),
        str(ExceptionType.CURRENCY_MISMATCH),
        str(ExceptionType.STATUS_INCONSISTENCY),
        str(ExceptionType.QA_INPUT_ERROR),
    }
